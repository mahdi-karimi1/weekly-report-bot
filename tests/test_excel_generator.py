"""Excel generator tests against the authoritative FRN template copy."""

from __future__ import annotations

import hashlib
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

from bot.models.report import TimedActivity, WeeklyReport
from bot.services.excel_generator import ExcelCapacityError, ExcelGenerator

TEMPLATE = Path("templates/weekly_report.xlsx")
MAPPING = Path("templates/weekly_report_mapping.json")
ORIGINAL_DESKTOP = Path(r"c:\Users\Asus\Desktop\FRN..Work-Report.xlsx")


def _file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


@pytest.fixture
def generator(tmp_path: Path) -> ExcelGenerator:
    assert TEMPLATE.exists()
    assert MAPPING.exists()
    return ExcelGenerator(TEMPLATE, MAPPING, tmp_path / "out")


def test_original_desktop_template_unchanged_by_generation(generator: ExcelGenerator) -> None:
    if not ORIGINAL_DESKTOP.exists():
        pytest.skip("Desktop original template not present")
    before = _file_sha256(ORIGINAL_DESKTOP)
    project_before = _file_sha256(TEMPLATE)

    report = WeeklyReport(
        employee_name="آزمایشی",
        report_date_range="۱ تا ۷ مهر",
        department_or_project="هوش مصنوعی",
        planned_completed=[TimedActivity(description="کار الف", hours=2)],
    )
    generator.generate(report, telegram_user_id=1, report_date=date(2026, 9, 26))

    assert _file_sha256(ORIGINAL_DESKTOP) == before
    assert _file_sha256(TEMPLATE) == project_before


def test_sample_data_cleared_and_missing_values_are_blank(generator: ExcelGenerator) -> None:
    report = WeeklyReport(
        employee_name="رضا",
        # intentionally omit report_date_range and department
        planned_completed=[TimedActivity(description="توسعه API", hours=4)],
        unplanned_completed=[],
        self_initiated=[],
        next_week_planned=[TimedActivity(description="تست یکپارچه")],  # hours missing
        problems_and_suggestions=["قطعی VPN"],
    )
    output = generator.generate(report, telegram_user_id=123, report_date=date(2026, 9, 26))

    wb = load_workbook(output)
    ws = wb["Sheet1"]

    # Sample template values must be gone.
    assert ws["B12"].value is None
    assert ws["C12"].value is None
    assert ws["H12"].value is None
    assert ws["J12"].value is None

    assert ws["B6"].value == "نام شخص: رضا"
    assert ws["F7"].value is None  # missing date -> blank
    assert ws["B7"].value is None  # missing department -> blank

    assert ws["B11"].value == "توسعه API"
    assert ws["C11"].value == 4
    assert ws["D11"].value is None
    assert ws["E11"].value is None
    assert ws["H11"].value == "تست یکپارچه"
    assert ws["I11"].value is None  # missing hours stay blank, not 0
    assert ws["J11"].value == "قطعی VPN"

    placeholders = {"—", "-", "نامشخص", "ذکر نشده", "N/A", "n/a"}
    for row in range(11, 26):
        for col in "BCDEFGHIJ":
            value = ws[f"{col}{row}"].value
            if isinstance(value, str):
                assert value.strip() not in placeholders
    # Explicitly ensure missing hours were not coerced to zero.
    assert ws["I11"].value is None

    wb.close()


def test_person_department_and_period_are_written_to_requested_cells(
    generator: ExcelGenerator,
) -> None:
    report = WeeklyReport(
        employee_name="نام کارمند",
        report_date_range="۱۴۰۵/۰۷/۰۴ الی ۱۴۰۵/۰۷/۰۸",
        department_or_project="فناوری",
    )

    output = generator.generate(report, telegram_user_id=44, report_date=date(2026, 9, 26))
    workbook = load_workbook(output)
    sheet = workbook["Sheet1"]

    assert sheet["B6"].value == "نام شخص: نام کارمند"
    assert sheet["B7"].value == "نام دپارتمان / پروژه: فناوری"
    assert sheet["F7"].value == "تاریخ گزارش: ۱۴۰۵/۰۷/۰۴ الی ۱۴۰۵/۰۷/۰۸"
    assert sheet["F6"].value is None
    workbook.close()


def test_activity_and_hours_adjacent_cells(generator: ExcelGenerator) -> None:
    report = WeeklyReport(
        planned_completed=[
            TimedActivity(description="A1", hours=1),
            TimedActivity(description="A2", hours=2.5),
        ],
        unplanned_completed=[TimedActivity(description="B1", hours=3)],
        self_initiated=[TimedActivity(description="C1")],
        next_week_planned=[TimedActivity(description="D1", hours=5)],
        problems_and_suggestions=["P1", "P2"],
    )
    output = generator.generate(report, telegram_user_id=9, report_date=date(2026, 9, 26))
    ws = load_workbook(output)["Sheet1"]
    assert ws["B11"].value == "A1" and ws["C11"].value == 1
    assert ws["B12"].value == "A2" and ws["C12"].value == 2.5
    assert ws["D11"].value == "B1" and ws["E11"].value == 3
    assert ws["F11"].value == "C1" and ws["G11"].value is None
    assert ws["H11"].value == "D1" and ws["I11"].value == 5
    assert ws["J11"].value == "P1" and ws["J12"].value == "P2"


def test_formatting_and_merged_cells_preserved(generator: ExcelGenerator) -> None:
    before = load_workbook(TEMPLATE)
    before_ws = before["Sheet1"]
    before_merges = {str(m) for m in before_ws.merged_cells.ranges}
    before_widths = {
        key: before_ws.column_dimensions[key].width
        for key in before_ws.column_dimensions
    }
    before_title = before_ws["D3"].value
    before_doc_no = before_ws["I2"].value
    before.close()

    report = WeeklyReport(planned_completed=[TimedActivity(description="x", hours=1)])
    output = generator.generate(report, telegram_user_id=5, report_date=date(2026, 9, 26))
    after = load_workbook(output)
    after_ws = after["Sheet1"]
    assert after.sheetnames == ["Sheet1"]
    assert {str(m) for m in after_ws.merged_cells.ranges} == before_merges
    assert after_ws["D3"].value == before_title
    assert after_ws["I2"].value == before_doc_no
    assert after_ws["I3"].value is not None  # revision label kept
    assert after_ws["I4"].value is not None  # publication date kept
    assert after_ws["B9"].value is not None  # section heading kept
    for key, width in before_widths.items():
        assert after_ws.column_dimensions[key].width == width
    after.close()


def test_overflow_beyond_15_rows_raises(generator: ExcelGenerator) -> None:
    activities = [TimedActivity(description=f"کار {i}", hours=1) for i in range(16)]
    report = WeeklyReport(planned_completed=activities)
    with pytest.raises(ExcelCapacityError) as exc:
        generator.generate(report, telegram_user_id=7, report_date=date(2026, 9, 26))
    assert exc.value.overflowing

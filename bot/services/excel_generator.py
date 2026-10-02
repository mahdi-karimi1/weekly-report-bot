"""Excel report generation from the authoritative FRN weekly template.

Loads a fresh copy of templates/weekly_report.xlsx for every report.
Clears only designated input cells, writes extracted values, and never
inserts placeholders for missing data.
"""

from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from bot.models.report import TEMPLATE_ACTIVITY_ROWS, TimedActivity, WeeklyReport
from bot.utils.logging import get_logger

logger = get_logger(__name__)

SAFE_FILENAME_RE = re.compile(r"[^A-Za-z0-9._-]+")
FORBIDDEN_PLACEHOLDERS = {
    "—",
    "-",
    "–",
    "نامشخص",
    "ذکر نشده",
    "n/a",
    "na",
    "none",
    "null",
}


class ExcelGenerationError(Exception):
    """Raised when Excel generation fails."""


class ExcelCapacityError(ExcelGenerationError):
    """Raised when a category exceeds the template's fixed row capacity."""

    def __init__(self, overflowing: list[str]) -> None:
        self.overflowing = overflowing
        categories = "، ".join(overflowing)
        super().__init__(
            f"Activity categories exceed template capacity (15 rows): {categories}"
        )


@dataclass
class FRNCellMapping:
    sheet_name: str = "Sheet1"
    person_name_cell: str = "B6"
    report_date_cell: str = "F7"
    department_project_cell: str = "B7"
    person_name_prefix: str = "نام شخص: "
    report_date_prefix: str = "تاریخ گزارش: "
    department_project_prefix: str = "نام دپارتمان / پروژه: "
    data_start_row: int = 11
    data_end_row: int = 25
    columns: dict[str, dict[str, str]] = field(default_factory=dict)
    clear_cells: list[str] = field(default_factory=list)
    clear_column_ranges: list[dict[str, Any]] = field(default_factory=list)

    @classmethod
    def from_dict(cls, data: dict[str, Any]) -> FRNCellMapping:
        return cls(
            sheet_name=data.get("sheet_name", "Sheet1"),
            person_name_cell=data.get("person_name_cell", "B6"),
            report_date_cell=data.get("report_date_cell", "F7"),
            department_project_cell=data.get("department_project_cell", "B7"),
            person_name_prefix=data.get("person_name_prefix", "نام شخص: "),
            report_date_prefix=data.get("report_date_prefix", "تاریخ گزارش: "),
            department_project_prefix=data.get(
                "department_project_prefix",
                "نام دپارتمان / پروژه: ",
            ),
            data_start_row=int(data.get("data_start_row", 11)),
            data_end_row=int(data.get("data_end_row", 25)),
            columns=data.get("columns", {}),
            clear_cells=list(data.get("clear_cells", [])),
            clear_column_ranges=list(data.get("clear_column_ranges", [])),
        )

    @property
    def capacity(self) -> int:
        return self.data_end_row - self.data_start_row + 1


class ExcelGenerator:
    def __init__(self, template_path: Path, mapping_path: Path, output_dir: Path) -> None:
        self.template_path = template_path
        self.mapping_path = mapping_path
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def load_mapping(self) -> FRNCellMapping:
        if not self.mapping_path.exists():
            raise ExcelGenerationError(f"Excel mapping file not found: {self.mapping_path}")
        data = json.loads(self.mapping_path.read_text(encoding="utf-8"))
        return FRNCellMapping.from_dict(data)

    def generate(
        self,
        report: WeeklyReport,
        *,
        telegram_user_id: int,
        report_date: date | None = None,
    ) -> Path:
        if not self.template_path.exists():
            raise ExcelGenerationError(
                f"Excel template not found: {self.template_path}. "
                "Copy FRN..Work-Report.xlsx to templates/weekly_report.xlsx."
            )

        mapping = self.load_mapping()
        overflowing = report.overflowing_categories(limit=mapping.capacity)
        if overflowing:
            raise ExcelCapacityError(overflowing)

        # Always work on a fresh load of the template copy — never the desktop original.
        workbook = load_workbook(self.template_path)
        if mapping.sheet_name not in workbook.sheetnames:
            raise ExcelGenerationError(
                f"Sheet '{mapping.sheet_name}' not found in template. "
                f"Available: {workbook.sheetnames}"
            )
        sheet = workbook[mapping.sheet_name]

        preserved_merges = {str(m) for m in sheet.merged_cells.ranges}
        preserved_title = workbook.sheetnames[:]

        self._clear_input_cells(sheet, mapping)
        self._write_header_fields(sheet, mapping, report)
        self._write_timed_column(
            sheet,
            mapping,
            mapping.columns.get("planned_completed", {}),
            report.planned_completed,
        )
        self._write_timed_column(
            sheet,
            mapping,
            mapping.columns.get("unplanned_completed", {}),
            report.unplanned_completed,
        )
        self._write_timed_column(
            sheet,
            mapping,
            mapping.columns.get("self_initiated", {}),
            report.self_initiated,
        )
        self._write_timed_column(
            sheet,
            mapping,
            mapping.columns.get("next_week_planned", {}),
            report.next_week_planned,
        )
        self._write_text_column(
            sheet,
            mapping,
            mapping.columns.get("problems_and_suggestions", {}).get("text"),
            report.problems_and_suggestions,
        )

        if {str(m) for m in sheet.merged_cells.ranges} != preserved_merges:
            workbook.close()
            raise ExcelGenerationError("Merged cells were altered during generation.")
        if workbook.sheetnames != preserved_title:
            workbook.close()
            raise ExcelGenerationError("Worksheet names were altered during generation.")

        output_path = self._build_output_path(telegram_user_id, report_date or date.today())
        workbook.save(output_path)
        workbook.close()
        logger.info("Generated Excel report at %s", output_path)
        return output_path

    def _build_output_path(self, telegram_user_id: int, report_date: date) -> Path:
        safe_user = SAFE_FILENAME_RE.sub("", str(telegram_user_id)) or "user"
        filename = f"weekly_report_{safe_user}_{report_date.isoformat()}.xlsx"
        return self.output_dir / filename

    def _clear_input_cells(self, sheet: Worksheet, mapping: FRNCellMapping) -> None:
        for cell_ref in mapping.clear_cells:
            sheet[cell_ref] = None
        for spec in mapping.clear_column_ranges:
            column = spec["column"]
            start_row = int(spec["start_row"])
            end_row = int(spec["end_row"])
            for row in range(start_row, end_row + 1):
                sheet[f"{column}{row}"] = None

    def _write_header_fields(
        self,
        sheet: Worksheet,
        mapping: FRNCellMapping,
        report: WeeklyReport,
    ) -> None:
        self._write_labeled_value(
            sheet,
            mapping.person_name_cell,
            mapping.person_name_prefix,
            report.employee_name,
        )
        self._write_labeled_value(
            sheet,
            mapping.report_date_cell,
            mapping.report_date_prefix,
            report.report_date_range,
        )
        self._write_labeled_value(
            sheet,
            mapping.department_project_cell,
            mapping.department_project_prefix,
            report.department_or_project,
        )

    def _write_labeled_value(
        self,
        sheet: Worksheet,
        cell_ref: str,
        prefix: str,
        value: str | None,
    ) -> None:
        cleaned = _sanitize_optional_text(value)
        if cleaned is None:
            sheet[cell_ref] = None
            return
        sheet[cell_ref] = f"{prefix}{cleaned}"

    def _write_timed_column(
        self,
        sheet: Worksheet,
        mapping: FRNCellMapping,
        columns: dict[str, str],
        activities: list[TimedActivity],
    ) -> None:
        text_col = columns.get("text")
        hours_col = columns.get("hours")
        if not text_col:
            return
        if len(activities) > mapping.capacity:
            raise ExcelCapacityError([text_col])

        for index, activity in enumerate(activities):
            row = mapping.data_start_row + index
            description = _sanitize_optional_text(activity.description)
            if description is None:
                continue
            sheet[f"{text_col}{row}"] = description
            if hours_col:
                hours = _sanitize_optional_hours(activity.hours)
                sheet[f"{hours_col}{row}"] = hours

    def _write_text_column(
        self,
        sheet: Worksheet,
        mapping: FRNCellMapping,
        column: str | None,
        values: list[str],
    ) -> None:
        if not column:
            return
        if len(values) > mapping.capacity:
            raise ExcelCapacityError([column])
        for index, value in enumerate(values):
            cleaned = _sanitize_optional_text(value)
            if cleaned is None:
                continue
            row = mapping.data_start_row + index
            sheet[f"{column}{row}"] = cleaned


def _sanitize_optional_text(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    if not cleaned:
        return None
    if cleaned.lower() in FORBIDDEN_PLACEHOLDERS or cleaned in FORBIDDEN_PLACEHOLDERS:
        return None
    return cleaned


def _sanitize_optional_hours(value: float | None) -> float | None:
    if value is None:
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    if number < 0:
        return None
    # Missing hours must stay blank — never coerce absent data to 0.
    return number


def install_frn_template(source: Path, destination: Path) -> Path:
    """Copy the authoritative FRN workbook into the project templates directory."""
    if not source.exists():
        raise FileNotFoundError(f"Source template not found: {source}")
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


# Compatibility stubs so older demo helpers do not break imports.
def create_demo_template(path: Path) -> Path:
    raise ExcelGenerationError(
        "Demo template creation is disabled. Use templates/weekly_report.xlsx "
        "(copy of FRN..Work-Report.xlsx)."
    )


def create_demo_mapping(path: Path) -> Path:
    raise ExcelGenerationError(
        "Demo mapping creation is disabled. Use templates/weekly_report_mapping.json."
    )
